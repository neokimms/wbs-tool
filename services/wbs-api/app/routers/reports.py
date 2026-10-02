"""대시보드·보고서·정기 보고 스케줄 API 라우터.
"""
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
from typing import Any
from uuid import UUID
import asyncpg
from fastapi import HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import APIRouter

from app.core.audit import insert_audit_event
from app.core.config import (
    APPROVAL_SELECT,
    AsyncIOScheduler,
    EXCEL_MEDIA_TYPE,
    REPORT_OUTPUT_DIR,
    REPORT_RUN_SELECT,
    REPORT_SCHEDULER_ENABLED,
    REPORT_SCHEDULE_SELECT,
    REPORT_SMTP_FROM,
    REPORT_SMTP_HOST,
    RESOURCE_DEFAULT_TASK_HOURS,
    logger,
)
from app.core.db import get_pool, normalize_metadata, normalize_record, parse_date, safe_uuid
from app.core.notify import normalize_email_recipients
from app.core.security import get_tenant_id, require_mutating_role, require_roles
from app.schemas import ReportRunRequest, ReportScheduleUpdate
from app.services.reports import (
    execute_report_schedule,
    fetch_report_schedule,
    refresh_report_scheduler,
    report_run_response,
    resource_identity_key,
    scheduler_runtime_state,
)
from app.services.wbs import fetch_rollups_for_projects

router = APIRouter()


def report_schedule_response(record: asyncpg.Record | dict[str, Any]) -> dict[str, Any]:
    schedule = normalize_record(record) if isinstance(record, asyncpg.Record) else dict(record)
    schedule["recipients"] = normalize_email_recipients(schedule.get("recipients"))
    schedule["smtp_configured"] = bool(REPORT_SMTP_HOST)
    schedule["scheduler_enabled"] = REPORT_SCHEDULER_ENABLED and AsyncIOScheduler is not None
    return schedule


@router.get("/api/report-schedules")
async def list_report_schedules(request: Request) -> dict[str, Any]:
    require_roles(request, {"admin", "pmo"})
    async with get_pool(request).acquire() as connection:
        records = await connection.fetch(
            f"""
            SELECT {REPORT_SCHEDULE_SELECT}
            FROM wbs_report_schedules
            ORDER BY key
            """
        )
    return {
        "schedules": [report_schedule_response(record) for record in records],
        "runtime": scheduler_runtime_state(request.app),
        "smtp": {
            "configured": bool(REPORT_SMTP_HOST),
            "host": REPORT_SMTP_HOST or None,
            "from": REPORT_SMTP_FROM,
        },
    }


@router.patch("/api/report-schedules/{schedule_key}")
async def update_report_schedule(
    schedule_key: str,
    payload: ReportScheduleUpdate,
    request: Request,
) -> dict[str, Any]:
    require_mutating_role(request)
    key = schedule_key.strip()
    if not key:
        raise HTTPException(status_code=400, detail="Schedule key is required")

    data = payload.model_dump(exclude_unset=True)
    recipients = normalize_email_recipients(data["recipients"]) if "recipients" in data else None
    metadata = data.get("metadata") if "metadata" in data else None

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            existing = await connection.fetchrow(
                f"""
                SELECT {REPORT_SCHEDULE_SELECT}
                FROM wbs_report_schedules
                WHERE key = $1
                FOR UPDATE
                """,
                key,
            )
            if not existing:
                raise HTTPException(status_code=404, detail="Report schedule not found")

            existing_metadata = normalize_metadata(existing["metadata"])
            next_metadata = {**existing_metadata, **(metadata or {})} if metadata is not None else existing_metadata
            record = await connection.fetchrow(
                f"""
                UPDATE wbs_report_schedules
                SET name = COALESCE($2, name),
                    enabled = COALESCE($3, enabled),
                    timezone = COALESCE($4, timezone),
                    day_of_week = COALESCE($5, day_of_week),
                    hour = COALESCE($6, hour),
                    minute = COALESCE($7, minute),
                    recipients = COALESCE($8::jsonb, recipients),
                    email_subject = COALESCE($9, email_subject),
                    email_body = COALESCE($10, email_body),
                    metadata = $11::jsonb,
                    updated_at = now()
                WHERE key = $1
                RETURNING {REPORT_SCHEDULE_SELECT}
                """,
                key,
                data.get("name"),
                data.get("enabled"),
                data.get("timezone"),
                data.get("day_of_week"),
                data.get("hour"),
                data.get("minute"),
                recipients,
                data.get("email_subject"),
                data.get("email_body"),
                next_metadata,
            )
            await insert_audit_event(
                connection,
                request=request,
                event_type="report.schedule_updated",
                entity_type="report_schedule",
                entity_id=key,
                summary=f"Report schedule updated: {key}",
                metadata={
                    "enabled": record["enabled"],
                    "day_of_week": record["day_of_week"],
                    "hour": record["hour"],
                    "minute": record["minute"],
                    "recipient_count": len(normalize_email_recipients(record["recipients"])),
                },
            )

    await refresh_report_scheduler(request.app)
    async with get_pool(request).acquire() as connection:
        refreshed = await fetch_report_schedule(connection, key)
    return {
        "schedule": report_schedule_response(refreshed),
        "runtime": scheduler_runtime_state(request.app),
    }


@router.post("/api/report-schedules/{schedule_key}/run", status_code=201)
async def run_report_schedule(
    schedule_key: str,
    request: Request,
    payload: ReportRunRequest = ReportRunRequest(),
) -> dict[str, Any]:
    require_mutating_role(request)
    async with get_pool(request).acquire() as connection:
        schedule = await fetch_report_schedule(connection, schedule_key.strip())
    if not schedule:
        raise HTTPException(status_code=404, detail="Report schedule not found")

    run = await execute_report_schedule(
        request.app,
        schedule,
        send_email=payload.send_email,
        triggered_by=payload.triggered_by,
        request=request,
    )
    await refresh_report_scheduler(request.app)
    return {
        "run": run,
        "runtime": scheduler_runtime_state(request.app),
    }


@router.get("/api/report-runs")
async def list_report_runs(request: Request, limit: int = 20) -> list[dict[str, Any]]:
    require_roles(request, {"admin", "pmo"})
    bounded_limit = min(max(limit, 1), 100)
    async with get_pool(request).acquire() as connection:
        records = await connection.fetch(
            f"""
            SELECT {REPORT_RUN_SELECT}
            FROM wbs_report_runs
            ORDER BY started_at DESC
            LIMIT $1
            """,
            bounded_limit,
        )
    return [report_run_response(record) for record in records]


@router.get("/api/report-runs/{run_id}/artifact")
async def download_report_run_artifact(run_id: str, request: Request) -> StreamingResponse:
    require_roles(request, {"admin", "pmo"})
    parsed_id = safe_uuid(run_id)
    if not parsed_id:
        raise HTTPException(status_code=400, detail="Invalid report run id")

    async with get_pool(request).acquire() as connection:
        record = await connection.fetchrow(
            f"""
            SELECT {REPORT_RUN_SELECT}
            FROM wbs_report_runs
            WHERE id = $1
            """,
            parsed_id,
        )
    if not record or not record["artifact_path"]:
        raise HTTPException(status_code=404, detail="Report artifact not found")

    artifact_path = Path(record["artifact_path"]).resolve()
    output_root = REPORT_OUTPUT_DIR.resolve()
    if output_root not in artifact_path.parents and artifact_path != output_root:
        raise HTTPException(status_code=403, detail="Report artifact path is outside the report directory")
    if not artifact_path.exists():
        raise HTTPException(status_code=404, detail="Report artifact file is missing")

    filename = artifact_path.name
    return StreamingResponse(
        BytesIO(artifact_path.read_bytes()),
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def _dashboard_phase_rows(request: Request, project_ids: list[UUID]) -> list[dict[str, Any]]:
    if not project_ids:
        return []
    async with get_pool(request).acquire() as connection:
        records = await connection.fetch(
            """
            SELECT i.project_id, i.code, i.name, i.parent_code, p.name AS project_name
            FROM wbs_project_wbs_items i
            JOIN wbs_projects p ON p.id = i.project_id
            WHERE i.project_id = ANY($1::uuid[]) AND i.item_type = '단계'
              AND (i.parent_code IS NULL OR NOT EXISTS (
                    SELECT 1 FROM wbs_project_wbs_items x
                    WHERE x.project_id = i.project_id AND x.code = i.parent_code AND x.item_type = '단계'))
            ORDER BY p.name, i.sort_order, i.code
            """,
            project_ids,
        )
    return [normalize_record(r) for r in records]


@router.get("/api/dashboard")
async def dashboard(request: Request) -> dict[str, Any]:
    tid = get_tenant_id(request)
    async with get_pool(request).acquire() as connection:
        project_count = await connection.fetchval("SELECT count(*) FROM wbs_projects WHERE tenant_id = $1", tid)
        template_count = await connection.fetchval("SELECT count(*) FROM wbs_templates")
        pending_approval_count = await connection.fetchval(
            "SELECT count(*) FROM wbs_approval_requests WHERE status = 'Pending' AND tenant_id = $1", tid
        )
        preview_import_count = await connection.fetchval(
            "SELECT count(*) FROM wbs_import_jobs WHERE status = 'Preview'"
        )
        status_rows = await connection.fetch(
            """
            SELECT status, count(*) AS count
            FROM wbs_projects
            WHERE tenant_id = $1
            GROUP BY status
            ORDER BY status
            """,
            tid,
        )
        latest_projects = await connection.fetch(
            """
            SELECT name, owner, status, start_date, template_key
            FROM wbs_projects
            WHERE tenant_id = $1
            ORDER BY created_at DESC
            LIMIT 5
            """,
            tid,
        )
        latest_approvals = await connection.fetch(
            f"""
            SELECT {APPROVAL_SELECT}
            FROM wbs_approval_requests a
            JOIN wbs_projects p ON p.id = a.project_id
            WHERE a.tenant_id = $1
            ORDER BY a.created_at DESC
            LIMIT 5
            """,
            tid,
        )

        # P2-01: 진척률·SPI — 공통 산식(wbs_rules) 기반 집계
        dash_projects = await connection.fetch(
            """
            SELECT p.id AS project_id, p.name AS project_name, p.status AS project_status
            FROM wbs_projects p
            WHERE p.status NOT IN ('Closed') AND p.tenant_id = $1
            """,
            tid,
        )
        dash_rollups = await fetch_rollups_for_projects(
            connection, [r["project_id"] for r in dash_projects]
        )

        # P2-01: 리스크·이슈 카운트 (wbs_risks / wbs_issues 테이블 존재 시)
        risk_count = 0
        issue_count = 0
        try:
            risk_count = await connection.fetchval(
                "SELECT count(*) FROM wbs_risks WHERE status != 'Closed' AND tenant_id = $1", tid
            ) or 0
            issue_count = await connection.fetchval(
                "SELECT count(*) FROM wbs_issues WHERE status != 'Closed' AND tenant_id = $1", tid
            ) or 0
        except Exception:
            logger.warning("대시보드 리스크/이슈 집계 실패", exc_info=True)

    # SPI(EV/PV)·진척률 — 공통 산식
    project_kpis = []
    heatmap_rows: list[dict[str, Any]] = []
    total_ev = total_pv = total_earned = total_weight_all = 0.0
    for r in sorted(dash_projects, key=lambda x: x["project_name"]):
        roll = dash_rollups.get(str(r["project_id"]), {"summary": {}, "nodes": {}})
        sm = roll["summary"]
        total_ev += sm.get("earned_value", 0.0)
        total_pv += sm.get("planned_value", 0.0)
        total_earned += sm.get("earned_weight", 0.0)
        total_weight_all += sm.get("total_weight", 0.0)
        project_kpis.append({
            "project_id": str(r["project_id"]),
            "project_name": r["project_name"],
            "project_status": r["project_status"],
            "total_weight": sm.get("total_weight", 0.0),
            "earned_weight": sm.get("earned_weight", 0.0),
            "progress_pct": sm.get("progress_pct", 0.0),
            "planned_pct": sm.get("planned_pct", 0.0),
            "spi": sm.get("spi"),
            "delayed_leaf_count": sm.get("delayed_leaf_count", 0),
        })
    # 단계별 히트맵: '단계' 유형 항목의 롤업 진척률 (프로젝트 최상위 단계 기준)
    phase_records = await _dashboard_phase_rows(request, [r["project_id"] for r in dash_projects])
    name_by_id = {str(r["project_id"]): r["project_name"] for r in dash_projects}
    for ph in phase_records:
        node = dash_rollups.get(str(ph["project_id"]), {"nodes": {}})["nodes"].get(ph["code"])
        if node is None:
            continue
        heatmap_rows.append({
            "project_name": name_by_id.get(str(ph["project_id"]), ""),
            "phase_name": ph["name"],
            "avg_progress": node["progress"],
        })
        if len(heatmap_rows) >= 50:
            break

    portfolio_spi = round(total_ev / total_pv, 3) if total_pv > 0 else None
    portfolio_progress = round(total_earned / total_weight_all * 100.0, 1) if total_weight_all else 0.0

    return {
        "metrics": {
            "projects": project_count,
            "templates": template_count,
            "pending_approvals": pending_approval_count,
            "preview_imports": preview_import_count,
            "risk_count": int(risk_count),
            "issue_count": int(issue_count),
            "portfolio_spi": portfolio_spi,
            "portfolio_progress": portfolio_progress,
            "progress_formula": "leaf 입력 · 상위 가중 롤업 / SPI = EV ÷ PV",
            "openproject_sync": "ready",
            "database": "PostgreSQL 17",
        },
        "status_distribution": [normalize_record(row) for row in status_rows],
        "latest_projects": [normalize_record(row) for row in latest_projects],
        "latest_approvals": [normalize_record(row) for row in latest_approvals],
        "project_kpis": project_kpis,
        "status_heatmap": [
            {
                "project_name": r["project_name"],
                "phase_name": r["phase_name"],
                "avg_progress": round(float(r["avg_progress"]), 1),
            }
            for r in heatmap_rows
        ],
        "risk_hotspots": [
            {"name": "Pending PMO approvals", "level": "attention" if pending_approval_count else "stable"},
            {"name": "Excel preview queue", "level": "watch" if preview_import_count else "stable"},
            {"name": "Open risks", "level": "attention" if risk_count > 0 else "stable", "count": int(risk_count)},
            {"name": "Open issues", "level": "watch" if issue_count > 0 else "stable", "count": int(issue_count)},
        ],
    }


@router.get("/api/reports/weekly-excel")
async def report_weekly_excel(request: Request, from_date: str | None = None, to_date: str | None = None) -> StreamingResponse:
    """PMO 주간 보고서 Excel — 프로젝트 현황·리스크·이슈·승인 이력 시트."""
    require_roles(request, {"admin", "pmo"})
    tid = get_tenant_id(request)
    today = date.today()
    dt_from = parse_date(from_date) if from_date else today - timedelta(days=7)
    dt_to   = parse_date(to_date)   if to_date   else today

    async with get_pool(request).acquire() as connection:
        projects = await connection.fetch(
            "SELECT * FROM wbs_projects WHERE tenant_id = $1 ORDER BY created_at DESC",
            tid,
        )
        approvals = await connection.fetch(
            """
            SELECT a.*, p.name AS project_name
            FROM wbs_approval_requests a
            JOIN wbs_projects p
              ON p.id = a.project_id
             AND p.tenant_id = a.tenant_id
            WHERE a.tenant_id = $1
              AND a.created_at::date BETWEEN $2 AND $3
            ORDER BY a.created_at DESC
            """,
            tid, dt_from, dt_to,
        )
        risks = await connection.fetch(
            """
            SELECT r.*, p.name AS project_name
            FROM wbs_risks r
            JOIN wbs_projects p
              ON p.id = r.project_id
             AND p.tenant_id = r.tenant_id
            WHERE r.tenant_id = $1
            ORDER BY r.created_at DESC
            """,
            tid,
        ) if await _table_exists(connection, "wbs_risks") else []
        issues = await connection.fetch(
            """
            SELECT i.*, p.name AS project_name
            FROM wbs_issues i
            JOIN wbs_projects p
              ON p.id = i.project_id
             AND p.tenant_id = i.tenant_id
            WHERE i.tenant_id = $1
            ORDER BY i.created_at DESC
            """,
            tid,
        ) if await _table_exists(connection, "wbs_issues") else []

    wb = Workbook()

    # ── 시트 1: 프로젝트 현황
    ws = wb.active
    ws.title = "프로젝트 현황"
    hdr_fill = PatternFill("solid", fgColor="0071E3")
    hdr_font = Font(bold=True, color="FFFFFF")
    headers = ["프로젝트명", "담당자", "유형", "상태", "시작일", "OpenProject ID"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    for r_idx, p in enumerate(projects, 2):
        meta = normalize_metadata(p.get("metadata") or {})
        ws.cell(row=r_idx, column=1, value=p["name"])
        ws.cell(row=r_idx, column=2, value=p["owner"])
        ws.cell(row=r_idx, column=3, value=p["template_key"])
        ws.cell(row=r_idx, column=4, value=p["status"])
        ws.cell(row=r_idx, column=5, value=str(p["start_date"]) if p.get("start_date") else "")
        ws.cell(row=r_idx, column=6, value=p.get("openproject_project_id") or "")
    for col in range(1, 7):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 22

    # ── 시트 2: 승인 이력
    ws2 = wb.create_sheet("승인 이력")
    hdrs2 = ["프로젝트", "제목", "요청 유형", "상태", "요청자", "검토자", "요청일", "결정일"]
    for col, h in enumerate(hdrs2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    for r_idx, a in enumerate(approvals, 2):
        ws2.cell(row=r_idx, column=1, value=a.get("project_name") or "")
        ws2.cell(row=r_idx, column=2, value=a["title"])
        ws2.cell(row=r_idx, column=3, value=a["request_type"])
        ws2.cell(row=r_idx, column=4, value=a["status"])
        ws2.cell(row=r_idx, column=5, value=a["requester"])
        ws2.cell(row=r_idx, column=6, value=a.get("reviewer") or "")
        ws2.cell(row=r_idx, column=7, value=str(a["created_at"].date()) if a.get("created_at") else "")
        ws2.cell(row=r_idx, column=8, value=str(a["decided_at"].date()) if a.get("decided_at") else "")
    for col in range(1, 9):
        ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = 20

    # ── 시트 3: 리스크
    ws3 = wb.create_sheet("리스크")
    hdrs3 = ["프로젝트", "제목", "심각도", "발생가능성", "담당자", "상태", "목표일", "대응전략"]
    for col, h in enumerate(hdrs3, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    for r_idx, r in enumerate(risks, 2):
        ws3.cell(row=r_idx, column=1, value=r.get("project_name") or "")
        ws3.cell(row=r_idx, column=2, value=r["title"])
        ws3.cell(row=r_idx, column=3, value=r["severity"])
        ws3.cell(row=r_idx, column=4, value=r["likelihood"])
        ws3.cell(row=r_idx, column=5, value=r["owner"])
        ws3.cell(row=r_idx, column=6, value=r["status"])
        ws3.cell(row=r_idx, column=7, value=str(r["due_date"]) if r.get("due_date") else "")
        ws3.cell(row=r_idx, column=8, value=r.get("mitigation") or "")
    for col in range(1, 9):
        ws3.column_dimensions[ws3.cell(row=1, column=col).column_letter].width = 20

    # ── 시트 4: 이슈
    ws4 = wb.create_sheet("이슈")
    hdrs4 = ["프로젝트", "제목", "우선순위", "담당자", "상태", "목표일", "설명"]
    for col, h in enumerate(hdrs4, 1):
        cell = ws4.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    for r_idx, i in enumerate(issues, 2):
        ws4.cell(row=r_idx, column=1, value=i.get("project_name") or "")
        ws4.cell(row=r_idx, column=2, value=i["title"])
        ws4.cell(row=r_idx, column=3, value=i["priority"])
        ws4.cell(row=r_idx, column=4, value=i["assignee"])
        ws4.cell(row=r_idx, column=5, value=i["status"])
        ws4.cell(row=r_idx, column=6, value=str(i["due_date"]) if i.get("due_date") else "")
        ws4.cell(row=r_idx, column=7, value=i.get("description") or "")
    for col in range(1, 8):
        ws4.column_dimensions[ws4.cell(row=1, column=col).column_letter].width = 20

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"PMO_Weekly_Report_{today.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf, media_type=EXCEL_MEDIA_TYPE,
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


async def _table_exists(connection: asyncpg.Connection, table_name: str) -> bool:
    result = await connection.fetchval(
        "SELECT EXISTS(SELECT 1 FROM information_schema.tables WHERE table_name=$1)", table_name
    )
    return bool(result)


@router.get("/api/reports/executive-summary")
async def report_executive_summary(request: Request) -> dict[str, Any]:
    """경영진 요약 — HTML 렌더링용 JSON (포털에서 PDF 인쇄)."""
    require_roles(request, {"admin", "pmo"})
    tid = get_tenant_id(request)
    async with get_pool(request).acquire() as connection:
        project_count = await connection.fetchval("SELECT count(*) FROM wbs_projects WHERE tenant_id = $1", tid)
        status_dist   = await connection.fetch(
            "SELECT status, count(*) AS cnt FROM wbs_projects WHERE tenant_id = $1 GROUP BY status ORDER BY cnt DESC",
            tid,
        )
        pending_count = await connection.fetchval(
            "SELECT count(*) FROM wbs_approval_requests WHERE tenant_id = $1 AND status = 'Pending'",
            tid,
        )
        risk_count = issue_count = 0
        if await _table_exists(connection, "wbs_risks"):
            risk_count  = await connection.fetchval(
                "SELECT count(*) FROM wbs_risks WHERE tenant_id = $1 AND status != 'Closed'",
                tid,
            ) or 0
            issue_count = await connection.fetchval(
                "SELECT count(*) FROM wbs_issues WHERE tenant_id = $1 AND status != 'Closed'",
                tid,
            ) or 0
        top_risks = []
        if await _table_exists(connection, "wbs_risks"):
            rows = await connection.fetch(
                """
                SELECT r.title, r.severity, r.status, p.name AS project_name
                FROM wbs_risks r
                JOIN wbs_projects p
                  ON p.id = r.project_id
                 AND p.tenant_id = r.tenant_id
                WHERE r.tenant_id = $1
                  AND r.status != 'Closed'
                  AND r.severity = '높음'
                ORDER BY r.created_at DESC
                LIMIT 5
                """,
                tid,
            )
            top_risks = [normalize_record(r) for r in rows]

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "summary": {
            "total_projects": project_count,
            "pending_approvals": pending_count,
            "open_risks": risk_count,
            "open_issues": issue_count,
        },
        "status_distribution": [{"status": r["status"], "count": r["cnt"]} for r in status_dist],
        "top_risks": top_risks,
    }


@router.get("/api/reports/burndown")
async def portfolio_burndown(
    request: Request,
    from_date: str | None = None,
    to_date: str | None = None,
    project_id: str | None = None,
) -> dict[str, Any]:
    """일별 earned weight 스냅샷 기반 번다운. 스냅샷 없으면 현재값 단일 포인트 반환."""
    tid = get_tenant_id(request)
    today = date.today()
    d_from = parse_date(from_date) or (today - timedelta(days=30))
    d_to = parse_date(to_date) or today
    async with get_pool(request).acquire() as connection:
        pid_filter = ""
        params: list[Any] = [tid, d_from, d_to]
        if project_id:
            p = safe_uuid(project_id)
            if p:
                params.append(p)
                pid_filter = f" AND s.project_id = ${len(params)}"

        snapshot_rows = await connection.fetch(
            f"""
            SELECT s.snapshot_date, SUM(s.earned_weight) AS earned, SUM(s.total_weight) AS total
            FROM wbs_progress_snapshots s
            JOIN wbs_projects p ON p.id = s.project_id
            WHERE p.tenant_id = $1 AND s.snapshot_date BETWEEN $2 AND $3{pid_filter}
            GROUP BY s.snapshot_date ORDER BY s.snapshot_date
            """,
            *params,
        ) if await _table_exists(connection, "wbs_progress_snapshots") else []

        # 스냅샷 없으면 현재 값(공통 산식)으로 단일 포인트 생성
        if not snapshot_rows:
            only_pid = safe_uuid(project_id) if project_id else None
            cur_ids = [r["id"] for r in await connection.fetch(
                "SELECT id FROM wbs_projects WHERE tenant_id = $1 AND status NOT IN ('Closed')"
                " AND ($2::uuid IS NULL OR id = $2::uuid)",
                tid, only_pid,
            )]
            cur_rollups = await fetch_rollups_for_projects(connection, cur_ids)
            earned = sum(r["summary"]["earned_weight"] for r in cur_rollups.values())
            total = sum(r["summary"]["total_weight"] for r in cur_rollups.values())
            snapshot_rows = [{"snapshot_date": today, "earned": earned, "total": total}]

    series = [
        {
            "date": str(r["snapshot_date"]),
            "earned": round(float(r["earned"] or 0), 2),
            "total": round(float(r["total"] or 0), 2),
            "progress_pct": round(float(r["earned"] or 0) / float(r["total"] or 1) * 100, 1),
        }
        for r in snapshot_rows
    ]
    return {"from_date": str(d_from), "to_date": str(d_to), "series": series}


def parse_effort_hours(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return max(float(value), 0.0)
    if isinstance(value, (int, float)):
        return max(float(value), 0.0)
    text = str(value).strip().lower().replace(",", ".")
    if not text:
        return 0.0
    match = re.search(r"(\d+(?:\.\d+)?)", text)
    return max(float(match.group(1)), 0.0) if match else 0.0


def count_weekdays(start: date, end: date) -> int:
    if end < start:
        start, end = end, start
    current = start
    total = 0
    while current <= end:
        if current.weekday() < 5:
            total += 1
        current += timedelta(days=1)
    return max(total, 1)


def utilization_state(value: float) -> str:
    if value >= 100:
        return "over"
    if value >= 80:
        return "high"
    if value >= 50:
        return "normal"
    return "low"


def empty_resource_bucket(name: str) -> dict[str, Any]:
    return {
        "assignee": name,
        "total_tasks": 0,
        "done_tasks": 0,
        "open_tasks": 0,
        "total_weight": 0.0,
        "total_effort_hours": 0.0,
        "estimated_effort_hours": 0.0,
        "inferred_effort_tasks": 0,
        "avg_progress": 0.0,
        "projects": [],
        "tasks": [],
        "_project_index": {},
        "_progress_sum": 0.0,
    }


def add_task_to_bucket(bucket: dict[str, Any], task: dict[str, Any]) -> None:
    bucket["total_tasks"] += 1
    bucket["done_tasks"] += 1 if task["done"] else 0
    bucket["open_tasks"] += 0 if task["done"] else 1
    bucket["total_weight"] += task["weight"]
    bucket["total_effort_hours"] += task["effort_hours"]
    bucket["estimated_effort_hours"] += task["estimated_effort_hours"]
    bucket["inferred_effort_tasks"] += 1 if task["effort_basis"] == "default" else 0
    bucket["_progress_sum"] += task["progress"]
    bucket["tasks"].append({
        "project_id": task["project_id"],
        "project_name": task["project_name"],
        "code": task["code"],
        "name": task["name"],
        "item_type": task["item_type"],
        "status": task["status"],
        "progress": task["progress"],
        "start_date": task["start_date"],
        "finish_date": task["finish_date"],
        "weight": task["weight"],
        "effort_hours": task["effort_hours"],
        "estimated_effort_hours": task["estimated_effort_hours"],
        "effort_basis": task["effort_basis"],
    })
    project = bucket["_project_index"].setdefault(task["project_id"], {
        "project_id": task["project_id"],
        "project_name": task["project_name"],
        "task_count": 0,
        "done_count": 0,
        "open_count": 0,
        "avg_progress": 0.0,
        "total_weight": 0.0,
        "total_effort_hours": 0.0,
        "estimated_effort_hours": 0.0,
        "_progress_sum": 0.0,
    })
    project["task_count"] += 1
    project["done_count"] += 1 if task["done"] else 0
    project["open_count"] += 0 if task["done"] else 1
    project["total_weight"] += task["weight"]
    project["total_effort_hours"] += task["effort_hours"]
    project["estimated_effort_hours"] += task["estimated_effort_hours"]
    project["_progress_sum"] += task["progress"]


def finalize_resource_bucket(bucket: dict[str, Any]) -> dict[str, Any]:
    total = max(int(bucket["total_tasks"]), 1)
    bucket["avg_progress"] = round(bucket["_progress_sum"] / total, 1)
    projects = []
    for project in bucket["_project_index"].values():
        project["avg_progress"] = round(project["_progress_sum"] / max(project["task_count"], 1), 1)
        project["total_weight"] = round(project["total_weight"], 2)
        project["total_effort_hours"] = round(project["total_effort_hours"], 1)
        project["estimated_effort_hours"] = round(project["estimated_effort_hours"], 1)
        project.pop("_progress_sum", None)
        projects.append(project)
    bucket["projects"] = sorted(projects, key=lambda item: (-item["task_count"], item["project_name"]))
    bucket["tasks"] = sorted(
        bucket["tasks"],
        key=lambda item: (item["finish_date"] or "9999-12-31", item["project_name"], item["code"] or ""),
    )
    bucket["total_weight"] = round(bucket["total_weight"], 2)
    bucket["total_effort_hours"] = round(bucket["total_effort_hours"], 1)
    bucket["estimated_effort_hours"] = round(bucket["estimated_effort_hours"], 1)
    bucket.pop("_project_index", None)
    bucket.pop("_progress_sum", None)
    return bucket

@router.get("/api/reports/resource-load")
async def resource_load(
    request: Request,
    from_date: str | None = None,
    to_date: str | None = None,
) -> dict[str, Any]:
    """담당자별 작업 부하, 사용자 계정 매핑, PMO 가동율을 함께 제공한다."""
    require_roles(request, {"admin", "pmo"})
    tid = get_tenant_id(request)
    today = date.today()
    d_from = parse_date(from_date) or date(today.year, today.month, 1)
    d_to = parse_date(to_date) or today
    if d_to < d_from:
        d_from, d_to = d_to, d_from
    workdays = count_weekdays(d_from, d_to)
    async with get_pool(request).acquire() as connection:
        users = await connection.fetch(
            """
            SELECT id, email, display_name, role, status, last_login_at
            FROM wbs_users
            WHERE tenant_id = $1
            ORDER BY CASE role WHEN 'admin' THEN 0 WHEN 'pmo' THEN 1 ELSE 2 END, email
            """,
            tid,
        )
        rows = await connection.fetch(
            """
            SELECT
                i.owner AS assignee,
                p.id AS project_id,
                p.name AS project_name,
                i.code,
                i.name AS task_name,
                i.item_type,
                i.start_date,
                i.finish_date,
                i.weight,
                i.metadata
            FROM wbs_project_wbs_items i
            JOIN wbs_projects p ON p.id = i.project_id
            WHERE p.tenant_id = $1
              AND i.item_type IN ('작업', '산출물', '마일스톤')
              AND i.owner IS NOT NULL AND i.owner <> ''
              AND (i.start_date IS NULL OR i.start_date <= $3)
              AND (i.finish_date IS NULL OR i.finish_date >= $2)
            ORDER BY i.owner, p.name, i.sort_order, i.code
            """,
            tid, d_from, d_to,
        )

    user_list = [normalize_record(user) for user in users]
    active_users = [user for user in user_list if user.get("status") == "Active"]
    pmo_users = [user for user in active_users if user.get("role") in {"admin", "pmo"}]
    fallback_pmo_user = next((user for user in active_users if user.get("role") == "pmo"), None)
    fallback_pmo_user = fallback_pmo_user or next((user for user in pmo_users), None)
    user_keys: dict[str, dict[str, Any]] = {}
    for user in active_users:
        candidates = [
            user.get("display_name"),
            user.get("email"),
            str(user.get("email") or "").partition("@")[0],
        ]
        for candidate in candidates:
            key = resource_identity_key(candidate)
            if key:
                user_keys.setdefault(key, user)

    def matched_user_for_owner(owner: str) -> dict[str, Any] | None:
        key = resource_identity_key(owner)
        if key in user_keys:
            return user_keys[key]
        if key == "pmo" and fallback_pmo_user:
            return fallback_pmo_user
        return None

    by_assignee: dict[str, dict[str, Any]] = {}
    account_buckets: dict[str, dict[str, Any]] = {}
    for user in active_users:
        uid = str(user["id"])
        account_buckets[uid] = {
            "user_id": uid,
            "email": user.get("email"),
            "display_name": user.get("display_name"),
            "role": user.get("role"),
            "status": user.get("status"),
            "last_login_at": user.get("last_login_at"),
            "matched_owners": [],
            **empty_resource_bucket(str(user.get("display_name") or user.get("email") or "")),
        }

    for r in rows:
        metadata = normalize_metadata(r["metadata"])
        status = str(metadata.get("status") or "")
        progress = parse_effort_hours(metadata.get("progress"))
        done = status == "완료" or progress >= 100
        effort_hours = parse_effort_hours(metadata.get("effort"))
        estimated_effort_hours = effort_hours if effort_hours > 0 else RESOURCE_DEFAULT_TASK_HOURS
        task = {
            "assignee": r["assignee"],
            "project_id": str(r["project_id"]),
            "project_name": r["project_name"],
            "code": r["code"],
            "name": r["task_name"],
            "item_type": r["item_type"],
            "status": status or "대기",
            "progress": min(progress, 100.0),
            "start_date": r["start_date"].isoformat() if r["start_date"] else None,
            "finish_date": r["finish_date"].isoformat() if r["finish_date"] else None,
            "weight": float(r["weight"] or 0),
            "effort_hours": effort_hours,
            "estimated_effort_hours": estimated_effort_hours,
            "effort_basis": "metadata" if effort_hours > 0 else "default",
            "done": done,
        }

        owner = task["assignee"]
        by_assignee.setdefault(owner, empty_resource_bucket(owner))
        matched_user = matched_user_for_owner(owner)
        if matched_user:
            by_assignee[owner]["matched_user"] = {
                "id": str(matched_user["id"]),
                "email": matched_user.get("email"),
                "display_name": matched_user.get("display_name"),
                "role": matched_user.get("role"),
            }
        add_task_to_bucket(by_assignee[owner], task)

        if matched_user:
            user_bucket = account_buckets[str(matched_user["id"])]
            if owner not in user_bucket["matched_owners"]:
                user_bucket["matched_owners"].append(owner)
            add_task_to_bucket(user_bucket, task)

    assignees = sorted(
        (finalize_resource_bucket(bucket) for bucket in by_assignee.values()),
        key=lambda item: (-item["total_tasks"], item["assignee"]),
    )
    account_tasks = sorted(
        (finalize_resource_bucket(bucket) for bucket in account_buckets.values()),
        key=lambda item: (-item["total_tasks"], item["role"], item["email"] or ""),
    )
    unmapped_assignments = [
        {
            "assignee": item["assignee"],
            "total_tasks": item["total_tasks"],
            "open_tasks": item["open_tasks"],
            "total_weight": item["total_weight"],
            "estimated_effort_hours": item["estimated_effort_hours"],
        }
        for item in assignees
        if not item.get("matched_user")
    ]
    account_by_user_id = {item["user_id"]: item for item in account_tasks}
    pmo_capacity = []
    capacity_hours = workdays * RESOURCE_DEFAULT_TASK_HOURS
    for user in pmo_users:
        uid = str(user["id"])
        account = account_by_user_id.get(uid) or {}
        planned_hours = float(account.get("estimated_effort_hours") or 0)
        utilization = round((planned_hours / capacity_hours) * 100, 1) if capacity_hours else 0.0
        pmo_capacity.append({
            "user_id": uid,
            "email": user.get("email"),
            "display_name": user.get("display_name"),
            "role": user.get("role"),
            "capacity_hours": round(capacity_hours, 1),
            "planned_hours": round(planned_hours, 1),
            "actual_effort_hours": round(float(account.get("total_effort_hours") or 0), 1),
            "inferred_effort_tasks": int(account.get("inferred_effort_tasks") or 0),
            "utilization_pct": utilization,
            "state": utilization_state(utilization),
            "task_count": int(account.get("total_tasks") or 0),
            "open_tasks": int(account.get("open_tasks") or 0),
        })
    summary = {
        "period_days": (d_to - d_from).days + 1,
        "workdays": workdays,
        "daily_capacity_hours": RESOURCE_DEFAULT_TASK_HOURS,
        "total_assignees": len(assignees),
        "total_tasks": sum(item["total_tasks"] for item in assignees),
        "done_tasks": sum(item["done_tasks"] for item in assignees),
        "open_tasks": sum(item["open_tasks"] for item in assignees),
        "total_weight": round(sum(float(item["total_weight"]) for item in assignees), 2),
        "total_effort_hours": round(sum(float(item["total_effort_hours"]) for item in assignees), 1),
        "estimated_effort_hours": round(sum(float(item["estimated_effort_hours"]) for item in assignees), 1),
        "mapped_accounts": sum(1 for item in account_tasks if item["total_tasks"] > 0),
        "unmapped_assignees": len(unmapped_assignments),
        "pmo_members": len(pmo_capacity),
        "pmo_capacity_hours": round(sum(item["capacity_hours"] for item in pmo_capacity), 1),
        "pmo_planned_hours": round(sum(item["planned_hours"] for item in pmo_capacity), 1),
    }
    return {
        "from_date": str(d_from),
        "to_date": str(d_to),
        "data_source": "internal_wbs",
        "summary": summary,
        "assignees": assignees,
        "account_tasks": account_tasks,
        "unmapped_assignments": unmapped_assignments,
        "pmo_capacity": sorted(pmo_capacity, key=lambda item: (-item["utilization_pct"], item["display_name"] or "")),
        "capacity_basis": {
            "daily_hours": RESOURCE_DEFAULT_TASK_HOURS,
            "workdays": workdays,
            "fallback": "공수 미입력 작업은 8h/건으로 추정",
        },
    }
