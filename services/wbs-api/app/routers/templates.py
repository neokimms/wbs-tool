"""표준 WBS 템플릿 API 라우터.
"""
from decimal import Decimal
from io import BytesIO
import re
from typing import Any
import asyncpg
from fastapi import File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.worksheet.datavalidation import DataValidation
from fastapi import APIRouter

from app.core.audit import insert_audit_event
from app.core.config import (
    AGILE_ITEM_TYPES,
    EXCEL_COLUMNS,
    EXCEL_MEDIA_TYPE,
    HEADER_ALIASES,
    IMPORT_JOB_RETURNING,
    MAX_EXCEL_UPLOAD_BYTES,
)
from app.core.db import as_date, get_pool, normalize_metadata, normalize_record, normalize_text
from app.core.security import require_mutating_role
from app.services.wbs import (
    assign_missing_wbs_codes,
    auto_code_warnings,
    build_wbs_diff_rows,
    fetch_template,
    fetch_template_items,
    infer_parent_from_code,
    insert_import_job,
    normalize_code,
    normalize_template_key,
    parent_depth,
    replace_template_items,
    root_code_from_rows,
    serialize_wbs_row,
    validate_wbs_rows,
)

router = APIRouter()


def normalize_header(value: Any) -> str:
    return re.sub(r"[\s_./-]+", "", str(value or "").strip().lower())


def as_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1]
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"Invalid number: {value}") from exc


def as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value or "").strip().lower()
    return text in {"1", "true", "t", "yes", "y", "예", "검수", "필수", "o"}


def template_code_prefix(template_key: str) -> str:
    known_prefixes = {
        "si-standard": "SI",
        "migration-data": "MIG",
        "maintenance": "OPS",
        "agile-standard": "AGL",
        "hybrid-standard": "HYB",
    }
    if template_key in known_prefixes:
        return known_prefixes[template_key]
    words = re.findall(r"[A-Z0-9]+", template_key.upper())
    if not words:
        return "WBS"
    return "".join(word[0] for word in words)[:6]


def template_delivery_mode(template: dict[str, Any]) -> str:
    key = str(template.get("key") or "").strip().lower()
    project_type = str(template.get("project_type") or "").strip().lower()
    if project_type == "agile" or key.startswith("agile"):
        return "agile"
    if project_type == "hybrid" or key.startswith("hybrid"):
        return "hybrid"
    return "waterfall"


def renumber_wbs_rows(rows: list[dict[str, Any]], root_code: str | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not rows:
        return [], []

    ordered_rows = sorted(rows, key=lambda row: (row.get("sort_order") or 0, row.get("code") or ""))
    children_by_parent: dict[str | None, list[dict[str, Any]]] = {}
    for row in ordered_rows:
        children_by_parent.setdefault(row.get("parent_code"), []).append(dict(row))

    roots = children_by_parent.get(None, [])
    if not roots:
        roots = [ordered_rows[0]]
        children_by_parent.setdefault(None, roots)

    resolved_root_code = root_code or roots[0].get("code") or "WBS"
    changes: list[dict[str, Any]] = []
    renumbered_rows: list[dict[str, Any]] = []

    def walk(row: dict[str, Any], new_code: str, new_parent_code: str | None) -> None:
        old_code = row.get("code")
        if old_code != new_code or row.get("parent_code") != new_parent_code:
            changes.append(
                {
                    "name": row.get("name"),
                    "old_code": old_code,
                    "new_code": new_code,
                    "old_parent_code": row.get("parent_code"),
                    "new_parent_code": new_parent_code,
                }
            )

        next_row = dict(row)
        next_row["code"] = new_code
        next_row["parent_code"] = new_parent_code
        renumbered_rows.append(next_row)

        counters = {"normal": 0, "milestone": 0}
        for child in children_by_parent.get(old_code, []):
            if child.get("item_type") == "마일스톤":
                counters["milestone"] += 1
                child_code = f"{new_code}.M{counters['milestone']}"
            else:
                counters["normal"] += 1
                child_code = f"{new_code}.{counters['normal']}"
            walk(child, child_code, new_code)

    for index, root in enumerate(roots, start=1):
        new_root_code = resolved_root_code if index == 1 else f"{resolved_root_code}.{index}"
        walk(root, new_root_code, None)

    for index, row in enumerate(renumbered_rows, start=1):
        row["sort_order"] = index

    return renumbered_rows, changes


def parse_wbs_workbook(contents: bytes) -> list[dict[str, Any]]:
    if len(contents) > MAX_EXCEL_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Excel file is too large")

    try:
        workbook = load_workbook(BytesIO(contents), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid Excel workbook: {exc}") from exc

    worksheet = workbook[workbook.sheetnames[0]]
    header_row = None
    header_map: dict[str, int] = {}

    for row_number in range(1, min(worksheet.max_row, 10) + 1):
        candidate: dict[str, int] = {}
        for column_number in range(1, worksheet.max_column + 1):
            header = HEADER_ALIASES.get(normalize_header(worksheet.cell(row_number, column_number).value))
            if header:
                candidate[header] = column_number
        if "name" in candidate:
            header_row = row_number
            header_map = candidate
            break

    if not header_row:
        raise HTTPException(status_code=400, detail="작업명 헤더를 찾을 수 없습니다")

    rows: list[dict[str, Any]] = []
    level_stack: dict[int, str] = {}

    for row_number in range(header_row + 1, worksheet.max_row + 1):
        raw = {
            field: worksheet.cell(row_number, column_number).value
            for field, column_number in header_map.items()
        }
        if not any(normalize_text(raw.get(field)) for field in ("code", "parent_code", "name")):
            continue

        try:
            level_value = as_float(raw.get("level"))
            level = int(level_value) if level_value else None
            code = normalize_code(raw.get("code"))
            parent_code = normalize_code(raw.get("parent_code"))
            if code and not parent_code and level and level > 1:
                parent_code = level_stack.get(level - 1)
            if code and not parent_code:
                parent_code = infer_parent_from_code(code)
            if code and level:
                level_stack[level] = code
                for stale_level in [key for key in level_stack if key > level]:
                    del level_stack[stale_level]

            rows.append(
                {
                    "row_number": row_number,
                    "level": level,
                    "code": code,
                    "name": normalize_text(raw.get("name")),
                    "parent_code": parent_code,
                    "item_type": normalize_text(raw.get("item_type")) or "작업",
                    "owner": normalize_text(raw.get("owner")),
                    "weight": as_float(raw.get("weight")),
                    "start_date": as_date(raw.get("start_date")),
                    "finish_date": as_date(raw.get("finish_date")),
                    "deliverable_type": normalize_text(raw.get("deliverable_type")),
                    "inspection_required": as_bool(raw.get("inspection_required")),
                    "progress_formula": normalize_text(raw.get("progress_formula")),
                    "notes": normalize_text(raw.get("notes")),
                }
            )
        except ValueError as exc:
            rows.append(
                {
                    "row_number": row_number,
                    "code": normalize_code(raw.get("code")),
                    "name": normalize_text(raw.get("name")),
                    "parent_code": normalize_code(raw.get("parent_code")),
                    "item_type": normalize_text(raw.get("item_type")) or "작업",
                    "parse_error": str(exc),
                }
            )

    return rows


def style_workbook_header(worksheet: Any, header_fill: PatternFill, header_font: Font, border: Border) -> None:
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def style_workbook_body(worksheet: Any, border: Border) -> None:
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_column_widths(worksheet: Any, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width


def add_agile_sample_sheets(
    workbook: Workbook,
    template: dict[str, Any],
    rows: list[dict[str, Any]],
    header_fill: PatternFill,
    header_font: Font,
    border: Border,
) -> None:
    delivery_mode = template_delivery_mode(template)
    if delivery_mode not in {"agile", "hybrid"}:
        return

    row_by_code = {row.get("code"): row for row in rows if row.get("code")}
    backlog = workbook.create_sheet("Agile Backlog")
    backlog.freeze_panes = "A2"
    backlog.append(
        [
            "유형",
            "제목",
            "설명",
            "Story Point",
            "우선순위",
            "상태",
            "담당자",
            "검토자",
            "Sprint",
            "연결 WBS 코드",
            "Acceptance Criteria",
            "Definition of Done",
        ]
    )
    style_workbook_header(backlog, header_fill, header_font, border)

    agile_row_count = 0
    for row in rows:
        metadata = normalize_metadata(row.get("metadata"))
        item_type = str(metadata.get("agile_type") or row.get("item_type") or "").strip()
        is_agile_item = item_type in AGILE_ITEM_TYPES or metadata.get("story_points") is not None
        if not is_agile_item:
            continue
        backlog.append(
            [
                item_type or "Task",
                row.get("name"),
                row.get("description") or metadata.get("notes"),
                metadata.get("story_points"),
                metadata.get("priority") or "Must",
                metadata.get("status") or "Backlog",
                row.get("owner"),
                metadata.get("reviewer"),
                metadata.get("sprint"),
                metadata.get("wbs_code") or (row.get("parent_code") if delivery_mode == "hybrid" else row.get("code")),
                metadata.get("acceptance_criteria"),
                metadata.get("definition_of_done"),
            ]
        )
        agile_row_count += 1

    style_workbook_body(backlog, border)
    set_column_widths(backlog, [14, 34, 32, 14, 14, 16, 16, 16, 18, 18, 44, 44])
    backlog.auto_filter.ref = backlog.dimensions
    type_validation = DataValidation(type="list", formula1='"Epic,Story,Task,Spike,Bug"', allow_blank=True)
    priority_validation = DataValidation(type="list", formula1='"Must,Should,Could,Wont"', allow_blank=True)
    status_validation = DataValidation(type="list", formula1='"Backlog,Ready,In Progress,Review,Done"', allow_blank=True)
    backlog.add_data_validation(type_validation)
    backlog.add_data_validation(priority_validation)
    backlog.add_data_validation(status_validation)
    type_validation.add("A2:A1000")
    priority_validation.add("E2:E1000")
    status_validation.add("F2:F1000")

    sprint_plan = workbook.create_sheet("Sprint Plan")
    sprint_plan.freeze_panes = "A2"
    sprint_plan.append(["Sprint", "목표", "시작일", "종료일", "계획 SP", "상태", "포함 항목"])
    style_workbook_header(sprint_plan, header_fill, header_font, border)

    sprints: dict[str, dict[str, Any]] = {}
    for row in rows:
        metadata = normalize_metadata(row.get("metadata"))
        sprint = metadata.get("sprint")
        if not sprint:
            continue
        entry = sprints.setdefault(
            str(sprint),
            {
                "goal": "",
                "start_date": None,
                "finish_date": None,
                "planned_points": 0,
                "planned_points_explicit": False,
                "status": "Planning",
                "items": [],
            },
        )
        if metadata.get("sprint_goal"):
            entry["goal"] = metadata.get("sprint_goal")
        if metadata.get("planned_points") is not None:
            entry["planned_points"] = metadata.get("planned_points")
            entry["planned_points_explicit"] = True
            entry["status"] = metadata.get("status") or entry["status"]
        if row.get("start_date") and not entry["start_date"]:
            entry["start_date"] = row.get("start_date")
        if row.get("finish_date") and not entry["finish_date"]:
            entry["finish_date"] = row.get("finish_date")
        if metadata.get("story_points") is not None:
            entry["items"].append(f"{row.get('name')} ({metadata.get('story_points')}SP)")
            if not entry.get("planned_points_explicit"):
                entry["planned_points"] = (entry.get("planned_points") or 0) + int(metadata.get("story_points") or 0)

    for sprint, entry in sprints.items():
        sprint_plan.append(
            [
                sprint,
                entry.get("goal"),
                entry.get("start_date"),
                entry.get("finish_date"),
                entry.get("planned_points"),
                entry.get("status"),
                ", ".join(entry.get("items") or []),
            ]
        )

    style_workbook_body(sprint_plan, border)
    set_column_widths(sprint_plan, [18, 46, 14, 14, 12, 16, 80])
    for row in sprint_plan.iter_rows(min_row=2):
        row[2].number_format = "yyyy-mm-dd"
        row[3].number_format = "yyyy-mm-dd"
    sprint_plan.auto_filter.ref = sprint_plan.dimensions

    if delivery_mode != "hybrid":
        return

    mapping = workbook.create_sheet("Hybrid Mapping")
    mapping.freeze_panes = "A2"
    mapping.append(["상위 WBS 코드", "상위 WBS 항목", "Agile 항목", "Sprint", "Story Point", "상태", "진행 기준"])
    style_workbook_header(mapping, header_fill, header_font, border)
    for row in rows:
        metadata = normalize_metadata(row.get("metadata"))
        if metadata.get("agile_type") not in AGILE_ITEM_TYPES and metadata.get("story_points") is None:
            continue
        wbs_code = metadata.get("wbs_code") or row.get("parent_code")
        parent = row_by_code.get(wbs_code)
        mapping.append(
            [
                wbs_code,
                parent.get("name") if parent else "",
                row.get("name"),
                metadata.get("sprint"),
                metadata.get("story_points"),
                metadata.get("status"),
                metadata.get("acceptance_criteria") or metadata.get("definition_of_done"),
            ]
        )
    style_workbook_body(mapping, border)
    set_column_widths(mapping, [18, 28, 34, 18, 14, 16, 56])
    mapping.auto_filter.ref = mapping.dimensions


def build_template_workbook(template: dict[str, Any], rows: list[dict[str, Any]]) -> BytesIO:
    parent_map = {row["code"]: row.get("parent_code") for row in rows if row.get("code")}
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WBS"
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="1D1D1F")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D8DCE2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    worksheet.append([label for _, label in EXCEL_COLUMNS])
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in rows:
        metadata = normalize_metadata(row.get("metadata"))
        start_date = row.get("start_date")
        finish_date = row.get("finish_date")
        worksheet.append(
            [
                parent_depth(row["code"], parent_map),
                row["code"],
                row.get("parent_code"),
                row["name"],
                row.get("item_type") or "작업",
                row.get("owner"),
                float(row["weight"]) if row.get("weight") is not None else None,
                start_date,
                finish_date,
                metadata.get("deliverable_type"),
                "Y" if metadata.get("inspection_required") else "N",
                metadata.get("progress_formula")
                or ("하위 단계 가중치 합산" if not row.get("parent_code") else "작업 완료율 x 가중치"),
                metadata.get("notes"),
            ]
        )

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[7].number_format = "yyyy-mm-dd"
        row[8].number_format = "yyyy-mm-dd"

    widths = [8, 18, 18, 30, 14, 16, 10, 14, 14, 18, 12, 24, 28]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width
    worksheet.auto_filter.ref = worksheet.dimensions

    item_type_validation = DataValidation(
        type="list",
        formula1='"프로젝트,단계,산출물,작업,마일스톤,리스크,이슈,변경요청"',
        allow_blank=False,
    )
    inspection_validation = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    worksheet.add_data_validation(item_type_validation)
    worksheet.add_data_validation(inspection_validation)
    item_type_validation.add("E2:E1000")
    inspection_validation.add("K2:K1000")

    guide = workbook.create_sheet("Guide")
    guide.append(["Template", template["name"]])
    guide.append(["Key", template["key"]])
    guide.append(["Project Type", template["project_type"]])
    guide.append(["Delivery Mode", template_delivery_mode(template)])
    guide.append(["Rule", "작업명은 필수입니다. WBS 코드는 비워두면 레벨과 행 순서 기준으로 자동 생성합니다."])
    if template_delivery_mode(template) in {"agile", "hybrid"}:
        guide.append(["Agile Sheets", "Agile Backlog와 Sprint Plan 시트에서 Epic/Story/Task, Story Point, Sprint 정보를 관리합니다."])
    if template_delivery_mode(template) == "hybrid":
        guide.append(["Hybrid Mapping", "상위 WBS 코드와 Agile 항목을 연결해 기준선 WBS와 Sprint 실행 현황을 함께 추적합니다."])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 96
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    add_agile_sample_sheets(workbook, template, rows, header_fill, header_font, border)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


async def prepare_template_import(
    connection: asyncpg.Connection,
    *,
    template_key: str,
    parsed_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    existing_rows = await fetch_template_items(connection, template_key)
    root_code = root_code_from_rows(existing_rows) or template_code_prefix(template_key)
    rows = assign_missing_wbs_codes(parsed_rows, root_code)
    errors, warnings = validate_wbs_rows(rows)
    warnings = [*warnings, *auto_code_warnings(rows)]
    serialized_rows = [serialize_wbs_row(row) for row in rows]
    diff_rows = build_wbs_diff_rows(existing_rows, rows)
    return rows, errors, warnings, serialized_rows, diff_rows


@router.get("/api/templates")
async def list_templates(request: Request) -> list[dict[str, Any]]:
    async with get_pool(request).acquire() as connection:
        records = await connection.fetch(
            """
            SELECT t.key, t.name, t.project_type, t.description, t.phases,
                   t.tenant_id,
                   count(i.id)::integer AS item_count
            FROM wbs_templates t
            LEFT JOIN wbs_template_items i ON i.template_key = t.key AND i.tenant_id = t.tenant_id
            GROUP BY t.key, t.name, t.project_type, t.description, t.phases, t.tenant_id
            ORDER BY t.project_type, t.name
            """
        )
    return [normalize_record(record) for record in records]


@router.get("/api/templates/{template_key}/items")
async def list_template_items(template_key: str, request: Request) -> dict[str, Any]:
    normalized_key = normalize_template_key(template_key)
    async with get_pool(request).acquire() as connection:
        template = await fetch_template(connection, normalized_key)
        if not template:
            raise HTTPException(status_code=404, detail="Template not found")
        rows = await fetch_template_items(connection, normalized_key)

    return {
        "template": template,
        "rows": rows,
    }


@router.get("/api/templates/{template_key}/versions")
async def list_template_versions(template_key: str, request: Request, limit: int = 20) -> list[dict[str, Any]]:
    normalized_key = normalize_template_key(template_key)
    bounded_limit = max(1, min(limit, 100))
    async with get_pool(request).acquire() as connection:
        records = await connection.fetch(
            """
            SELECT id, template_key, version, template_name, project_type,
                   description, item_count, metadata, created_at
            FROM wbs_template_versions
            WHERE template_key = $1
            ORDER BY version DESC
            LIMIT $2
            """,
            normalized_key,
            bounded_limit,
        )
    return [normalize_record(record) for record in records]


@router.get("/api/templates/{template_key}/excel")
async def export_template_excel(template_key: str, request: Request) -> StreamingResponse:
    normalized_key = normalize_template_key(template_key)
    async with get_pool(request).acquire() as connection:
        template = await fetch_template(connection, normalized_key)
        if not template:
            raise HTTPException(status_code=404, detail="Template not found")
        rows = await fetch_template_items(connection, normalized_key)

    if not rows:
        rows = [
            {
                "code": phase["code"],
                "parent_code": None,
                "name": phase["name"],
                "item_type": "단계",
                "owner": "PMO",
                "weight": phase.get("weight"),
                "start_date": None,
                "finish_date": None,
                "metadata": {},
            }
            for phase in template.get("phases", [])
        ]

    output = build_template_workbook(template, rows)
    filename = f"{normalized_key}-wbs-template.xlsx"
    return StreamingResponse(
        output,
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/api/templates/import/preview", status_code=201)
async def preview_template_excel(
    request: Request,
    file: UploadFile = File(...),
    template_key: str = Form("uploaded-wbs"),
    template_name: str = Form("업로드 WBS 템플릿"),
    project_type: str = Form("Uploaded"),
    description: str = Form("Excel 업로드로 반영될 계층형 WBS 템플릿"),
) -> dict[str, Any]:
    require_mutating_role(request)
    normalized_key = normalize_template_key(template_key)
    contents = await file.read()
    parsed_rows = parse_wbs_workbook(contents)

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            rows, errors, warnings, serialized_rows, diff_rows = await prepare_template_import(
                connection,
                template_key=normalized_key,
                parsed_rows=parsed_rows,
            )
            status = "Rejected" if errors else "Preview"
            record = await insert_import_job(
                connection,
                source_file=file.filename or "wbs-upload.xlsx",
                template_key=normalized_key,
                template_name=template_name,
                project_type=project_type,
                description=description,
                status=status,
                total_rows=len(rows),
                accepted_rows=0 if errors else len(rows),
                rejected_rows=len(errors),
                errors=errors,
                warnings=warnings,
                preview_rows=serialized_rows,
                diff_rows=diff_rows,
            )
            await insert_audit_event(
                connection,
                request=request,
                event_type="import.previewed",
                entity_type="import_job",
                entity_id=record["id"],
                summary=f"Excel import preview: {file.filename or 'wbs-upload.xlsx'}",
                metadata={
                    "template_key": normalized_key,
                    "status": status,
                    "accepted_rows": 0 if errors else len(rows),
                    "rejected_rows": len(errors),
                    "diff_count": len(diff_rows),
                },
            )
            template = await fetch_template(connection, normalized_key)

    response = normalize_record(record)
    response["template"] = template
    response["warnings"] = warnings
    response["rows"] = serialized_rows[:50]
    response["diff_rows"] = diff_rows[:50]
    return response


@router.post("/api/templates/import", status_code=201)
async def import_template_excel(
    request: Request,
    file: UploadFile = File(...),
    template_key: str = Form("uploaded-wbs"),
    template_name: str = Form("업로드 WBS 템플릿"),
    project_type: str = Form("Uploaded"),
    description: str = Form("Excel 업로드로 반영된 계층형 WBS 템플릿"),
) -> dict[str, Any]:
    require_mutating_role(request)
    normalized_key = normalize_template_key(template_key)
    contents = await file.read()
    parsed_rows = parse_wbs_workbook(contents)

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            rows, errors, warnings, serialized_rows, diff_rows = await prepare_template_import(
                connection,
                template_key=normalized_key,
                parsed_rows=parsed_rows,
            )
            status = "Rejected" if errors else "Accepted"
            record = await insert_import_job(
                connection,
                source_file=file.filename or "wbs-upload.xlsx",
                template_key=normalized_key,
                template_name=template_name,
                project_type=project_type,
                description=description,
                status=status,
                total_rows=len(rows),
                accepted_rows=0 if errors else len(rows),
                rejected_rows=len(errors),
                errors=errors,
                warnings=warnings,
                preview_rows=serialized_rows,
                diff_rows=diff_rows,
                applied=not errors,
            )

            if not errors:
                await replace_template_items(
                    connection,
                    template_key=normalized_key,
                    template_name=template_name,
                    project_type=project_type,
                    description=description,
                    rows=rows,
                )
                template_version = await connection.fetchval(
                    "SELECT max(version) FROM wbs_template_versions WHERE template_key = $1",
                    normalized_key,
                )
                record = await connection.fetchrow(
                    f"""
                    UPDATE wbs_import_jobs
                    SET template_version = $2
                    WHERE id = $1
                    RETURNING {IMPORT_JOB_RETURNING}
                    """,
                    record["id"],
                    template_version,
                )
            await insert_audit_event(
                connection,
                request=request,
                event_type="import.created",
                entity_type="import_job",
                entity_id=record["id"],
                summary=f"Excel import {'accepted' if not errors else 'rejected'}: {file.filename or 'wbs-upload.xlsx'}",
                metadata={
                    "template_key": normalized_key,
                    "status": status,
                    "accepted_rows": 0 if errors else len(rows),
                    "rejected_rows": len(errors),
                    "diff_count": len(diff_rows),
                    "applied": not errors,
                },
            )

            template = await fetch_template(connection, normalized_key)

    response = normalize_record(record)
    response["template"] = template
    response["warnings"] = warnings
    response["rows"] = serialized_rows[:50]
    response["diff_rows"] = diff_rows[:50]
    return response


@router.post("/api/templates/{template_key}/codes/resequence")
async def resequence_template_codes(template_key: str, request: Request) -> dict[str, Any]:
    require_mutating_role(request)
    normalized_key = normalize_template_key(template_key)

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            template = await fetch_template(connection, normalized_key)
            if not template:
                raise HTTPException(status_code=404, detail="Template not found")

            rows = await fetch_template_items(connection, normalized_key)
            if not rows:
                raise HTTPException(status_code=404, detail="Template has no WBS rows")

            root_code = root_code_from_rows(rows) or template_code_prefix(normalized_key)
            renumbered_rows, changes = renumber_wbs_rows(rows, root_code)
            await replace_template_items(
                connection,
                template_key=normalized_key,
                template_name=template["name"],
                project_type=template["project_type"],
                description=template["description"],
                rows=renumbered_rows,
            )
            updated_template = await fetch_template(connection, normalized_key)
            await insert_audit_event(
                connection,
                request=request,
                event_type="template.resequenced",
                entity_type="template",
                entity_id=normalized_key,
                summary=f"WBS codes resequenced: {normalized_key}",
                metadata={"changed_rows": len(changes)},
            )

    return {
        "template": updated_template,
        "status": "Updated",
        "changed_rows": len(changes),
        "changes": changes[:50],
        "rows": [serialize_wbs_row(row) for row in renumbered_rows[:50]],
    }
