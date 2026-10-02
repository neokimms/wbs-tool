"""Excel 가져오기 API 라우터.
"""
from datetime import date
from io import BytesIO
from typing import Any
from uuid import UUID
import asyncpg
from fastapi import HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import APIRouter

from app.core.audit import insert_audit_event
from app.core.config import EXCEL_COLUMNS, EXCEL_MEDIA_TYPE, IMPORT_JOB_RETURNING, PROJECT_WBS_IMPORT_ALLOWED_STATUSES
from app.core.db import get_pool, normalize_record, safe_uuid
from app.core.security import ensure_project_mutate_role, require_mutating_role
from app.schemas import WbsImportValidation
from app.services.projects import ensure_project_status_allowed, fetch_tenant_project
from app.services.wbs import (
    assign_missing_wbs_codes,
    auto_code_warnings,
    build_wbs_diff_rows,
    fetch_project_wbs_items,
    fetch_template,
    fetch_template_items,
    insert_import_job,
    normalize_template_key,
    replace_project_wbs_items,
    replace_template_items,
    serialize_wbs_row,
    validate_wbs_rows,
)

router = APIRouter()


def restore_wbs_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    restored_rows: list[dict[str, Any]] = []
    for row in rows:
        restored = dict(row)
        for key in ("start_date", "finish_date"):
            value = restored.get(key)
            if isinstance(value, str) and value:
                restored[key] = date.fromisoformat(value)
        restored_rows.append(restored)
    return restored_rows


def build_import_errors_workbook(job: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    issues = workbook.active
    issues.title = "Issues"
    issues.append(["Type", "Row", "Field", "Code", "Message"])
    for issue_type, entries in (("Error", job.get("errors") or []), ("Warning", job.get("warnings") or [])):
        for issue in entries:
            issues.append(
                [
                    issue_type,
                    issue.get("row"),
                    issue.get("field") or issue.get("parent_code"),
                    issue.get("code") or issue.get("parent_code"),
                    issue.get("message"),
                ]
            )

    diff = workbook.create_sheet("Diff")
    diff.append(["Change", "Code", "Name", "Field", "Before", "After"])
    for item in job.get("diff_rows") or []:
        fields = item.get("fields") or [{}]
        for field in fields:
            diff.append(
                [
                    item.get("change"),
                    item.get("code"),
                    item.get("name"),
                    field.get("field"),
                    field.get("before"),
                    field.get("after"),
                ]
            )

    rows = workbook.create_sheet("Rows")
    rows.append([label for _, label in EXCEL_COLUMNS])
    for row in job.get("preview_rows") or []:
        rows.append([row.get(key) for key, _ in EXCEL_COLUMNS])

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="1D1D1F")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column in range(1, worksheet.max_column + 1):
            worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = 22

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


async def fetch_import_job_for_update(connection: asyncpg.Connection, job_id: UUID) -> dict[str, Any] | None:
    record = await connection.fetchrow(
        f"""
        SELECT {IMPORT_JOB_RETURNING}
        FROM wbs_import_jobs
        WHERE id = $1
        FOR UPDATE
        """,
        job_id,
    )
    return normalize_record(record) if record else None


def import_job_response(record: asyncpg.Record, *, include_rows: bool = False, row_limit: int = 50) -> dict[str, Any]:
    job = normalize_record(record)
    preview_rows = job.get("preview_rows") or []
    diff_rows = job.get("diff_rows") or []
    job["preview_count"] = len(preview_rows)
    job["diff_count"] = len(diff_rows)
    if include_rows:
        job["rows"] = preview_rows[:row_limit]
        job["diff_rows"] = diff_rows[:row_limit]
    job.pop("preview_rows", None)
    return job


@router.get("/api/imports")
async def list_import_jobs(
    request: Request,
    limit: int = 10,
    status: str | None = None,
) -> list[dict[str, Any]]:
    bounded_limit = max(1, min(limit, 50))

    async with get_pool(request).acquire() as connection:
        if status:
            records = await connection.fetch(
                f"""
                SELECT {IMPORT_JOB_RETURNING}
                FROM wbs_import_jobs
                WHERE status = $1
                ORDER BY created_at DESC
                LIMIT $2
                """,
                status,
                bounded_limit,
            )
        else:
            records = await connection.fetch(
                f"""
                SELECT {IMPORT_JOB_RETURNING}
                FROM wbs_import_jobs
                ORDER BY created_at DESC
                LIMIT $1
                """,
                bounded_limit,
            )

    return [import_job_response(record) for record in records]


@router.get("/api/imports/{job_id}")
async def get_import_job(job_id: str, request: Request) -> dict[str, Any]:
    try:
        import_job_id = UUID(job_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid import job id") from exc

    async with get_pool(request).acquire() as connection:
        record = await connection.fetchrow(
            f"""
            SELECT {IMPORT_JOB_RETURNING}
            FROM wbs_import_jobs
            WHERE id = $1
            """,
            import_job_id,
        )

    if not record:
        raise HTTPException(status_code=404, detail="Import job not found")
    return import_job_response(record, include_rows=True)


@router.get("/api/imports/{job_id}/errors.xlsx")
async def download_import_errors(job_id: str, request: Request) -> StreamingResponse:
    try:
        import_job_id = UUID(job_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid import job id") from exc

    async with get_pool(request).acquire() as connection:
        record = await connection.fetchrow(
            f"""
            SELECT {IMPORT_JOB_RETURNING}
            FROM wbs_import_jobs
            WHERE id = $1
            """,
            import_job_id,
        )

    if not record:
        raise HTTPException(status_code=404, detail="Import job not found")
    job = normalize_record(record)
    output = build_import_errors_workbook(job)
    filename = f"{job['source_file'].rsplit('.', 1)[0]}-issues.xlsx"
    return StreamingResponse(
        output,
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/api/imports/validate")
async def validate_import(payload: WbsImportValidation, request: Request) -> dict[str, Any]:
    require_mutating_role(request)
    rows = assign_missing_wbs_codes([
        {"row_number": index, **row.model_dump()}
        for index, row in enumerate(payload.rows, start=1)
    ], "WBS")
    errors, warnings = validate_wbs_rows(rows)
    warnings = [*warnings, *auto_code_warnings(rows)]
    serialized_rows = [serialize_wbs_row(row) for row in rows]

    status = "Rejected" if errors else "Accepted"
    async with get_pool(request).acquire() as connection:
        record = await insert_import_job(
            connection,
            source_file=payload.source_file,
            status=status,
            total_rows=len(rows),
            accepted_rows=0 if errors else len(rows),
            rejected_rows=len(errors),
            errors=errors,
            warnings=warnings,
            preview_rows=serialized_rows,
        )

    response = normalize_record(record)
    response["warnings"] = warnings
    response["rows"] = serialized_rows
    return response


@router.delete("/api/imports/{job_id}", status_code=204)
async def delete_import_job(job_id: str, request: Request) -> None:
    require_mutating_role(request)
    parsed_id = safe_uuid(job_id)
    if not parsed_id:
        raise HTTPException(status_code=400, detail="Invalid import job id")

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            job = await connection.fetchrow(
                "SELECT id, status, source_file FROM wbs_import_jobs WHERE id = $1",
                parsed_id,
            )
            if not job:
                raise HTTPException(status_code=404, detail="Import job not found")
            if job["status"] == "Applied":
                raise HTTPException(
                    status_code=409,
                    detail="Applied 상태의 import job은 삭제할 수 없습니다. 먼저 프로젝트 WBS를 초기화하세요.",
                )
            await connection.execute("DELETE FROM wbs_import_jobs WHERE id = $1", parsed_id)
            await insert_audit_event(
                connection,
                request=request,
                event_type="import.deleted",
                entity_type="import_job",
                entity_id=parsed_id,
                summary=f"Import job deleted: {job['source_file']}",
                metadata={"status": job["status"]},
            )


@router.post("/api/imports/{job_id}/apply")
async def apply_import_preview(job_id: str, request: Request) -> dict[str, Any]:
    require_mutating_role(request)
    try:
        import_job_id = UUID(job_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid import job id") from exc

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            job = await fetch_import_job_for_update(connection, import_job_id)
            if not job:
                raise HTTPException(status_code=404, detail="Import job not found")
            if job["status"] != "Preview":
                raise HTTPException(status_code=409, detail="Import job is not waiting for approval")
            if job.get("errors"):
                raise HTTPException(status_code=400, detail="Rejected import cannot be applied")
            if not job.get("template_key"):
                raise HTTPException(status_code=400, detail="Import job has no template key")

            rows = restore_wbs_rows(job.get("preview_rows") or [])
            errors, _ = validate_wbs_rows(rows)
            if errors:
                raise HTTPException(status_code=400, detail={"message": "Preview rows are no longer valid", "errors": errors})

            template_key = normalize_template_key(job["template_key"])
            await replace_template_items(
                connection,
                template_key=template_key,
                template_name=job.get("template_name") or "업로드 WBS 템플릿",
                project_type=job.get("project_type") or "Uploaded",
                description=job.get("description") or "Excel 업로드로 반영된 계층형 WBS 템플릿",
                rows=rows,
            )
            template_version = await connection.fetchval(
                "SELECT max(version) FROM wbs_template_versions WHERE template_key = $1",
                template_key,
            )
            record = await connection.fetchrow(
                f"""
                UPDATE wbs_import_jobs
                SET status = 'Applied',
                    accepted_rows = total_rows,
                    rejected_rows = 0,
                    template_version = $2,
                    applied_at = now()
                WHERE id = $1
                RETURNING {IMPORT_JOB_RETURNING}
                """,
                import_job_id,
                template_version,
            )
            template = await fetch_template(connection, template_key)
            await insert_audit_event(
                connection,
                request=request,
                event_type="import.applied",
                entity_type="import_job",
                entity_id=import_job_id,
                summary=f"Excel import applied: {job['source_file']}",
                metadata={
                    "template_key": template_key,
                    "accepted_rows": record["accepted_rows"],
                },
            )

    response = normalize_record(record)
    response["template"] = template
    response["warnings"] = response.get("warnings", [])
    response["rows"] = (response.get("preview_rows") or [])[:50]
    return response


@router.post("/api/projects/{project_id}/imports/{job_id}/apply")
async def apply_import_to_project(project_id: str, job_id: str, request: Request) -> dict[str, Any]:
    parsed_project_id = safe_uuid(project_id)
    import_job_id = safe_uuid(job_id)
    if not parsed_project_id:
        raise HTTPException(status_code=400, detail="Invalid project id")
    if not import_job_id:
        raise HTTPException(status_code=400, detail="Invalid import job id")

    async with get_pool(request).acquire() as connection:
        async with connection.transaction():
            project = await fetch_tenant_project(connection, request, parsed_project_id, for_update=True)
            ensure_project_mutate_role(project)
            ensure_project_status_allowed(project, PROJECT_WBS_IMPORT_ALLOWED_STATUSES, "Project WBS import")

            job = await fetch_import_job_for_update(connection, import_job_id)
            if not job:
                raise HTTPException(status_code=404, detail="Import job not found")
            if job["status"] not in {"Preview", "Applied", "Accepted"}:
                raise HTTPException(status_code=409, detail="Import job is not valid for project apply")
            if job.get("errors"):
                raise HTTPException(status_code=400, detail="Rejected import cannot be applied to project")

            rows = restore_wbs_rows(job.get("preview_rows") or [])
            errors, _ = validate_wbs_rows(rows)
            if errors:
                raise HTTPException(status_code=400, detail={"message": "Import rows are no longer valid", "errors": errors})

            existing_rows = await fetch_project_wbs_items(connection, parsed_project_id)
            if not existing_rows:
                existing_rows = await fetch_template_items(connection, project["template_key"])
            diff_rows = build_wbs_diff_rows(existing_rows, rows)
            await replace_project_wbs_items(
                connection,
                project_id=parsed_project_id,
                tenant_id=project["tenant_id"],
                rows=rows,
                source_import_job_id=import_job_id,
            )
            record = await connection.fetchrow(
                """
                UPDATE wbs_projects
                SET status = CASE WHEN status = 'Rejected' THEN 'Draft' ELSE status END,
                    updated_at = now()
                WHERE id = $1
                RETURNING id, name, template_key, owner, status, start_date, delivery_mode,
                          openproject_project_id, metadata, tenant_id, created_at, updated_at
                """,
                parsed_project_id,
            )
            await insert_audit_event(
                connection,
                request=request,
                event_type="project_wbs.import_applied",
                entity_type="project",
                entity_id=parsed_project_id,
                summary=f"Project WBS import applied: {project['name']}",
                metadata={
                    "import_job_id": str(import_job_id),
                    "source_file": job.get("source_file"),
                    "rows": len(rows),
                    "diff_count": len(diff_rows),
                },
            )

    return {
        "status": "Applied",
        "project": normalize_record(record),
        "rows": [serialize_wbs_row(row) for row in rows[:50]],
        "diff_rows": diff_rows[:50],
        "summary": {"rows": len(rows), "diff_count": len(diff_rows)},
    }
