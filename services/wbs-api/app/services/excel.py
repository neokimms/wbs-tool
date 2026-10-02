"""Excel 보고서 작성 유틸리티.
"""
from datetime import date, datetime
from io import BytesIO
from typing import Any
import asyncpg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

from app.core.config import APPROVAL_SELECT, SYNC_RUN_SELECT
from app.services.wbs import fetch_rollups_for_projects


def worksheet_timestamp(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def append_report_table(worksheet, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    if worksheet.max_row > 1 or worksheet.cell(row=1, column=1).value:
        worksheet.append([])
    worksheet.append([title])
    title_row = worksheet.max_row
    worksheet.cell(title_row, 1).font = Font(bold=True, size=13)
    worksheet.append(headers)
    header_row = worksheet.max_row
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    thin_border = Border(bottom=Side(style="thin", color="D7DFEA"))
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        worksheet.append([worksheet_timestamp(value) for value in row])


def autosize_worksheet(worksheet) -> None:
    for column_cells in worksheet.columns:
        length = 0
        column = column_cells[0].column_letter
        for cell in column_cells:
            length = max(length, len(str(cell.value or "")))
        worksheet.column_dimensions[column].width = min(max(length + 2, 12), 42)


# ── 주간 프로젝트 보고서 ─────────────────────────────────────────────────
# 데이터 수집(_fetch_weekly_report_data) → 시트 정의(_weekly_report_sheets) → 작성(build_weekly_report_workbook)
# 로 분리. 진척 수치는 공통 산식(compute_wbs_rollup) 결과만 사용한다.

async def _fetch_weekly_report_data(
    pool: asyncpg.Pool, period_start: datetime, period_end: datetime
) -> dict[str, Any]:
    async with pool.acquire() as connection:
        counts = {
            key: await connection.fetchval(sql)
            for key, sql in (
                ("project_count", "SELECT count(*) FROM wbs_projects"),
                ("template_count", "SELECT count(*) FROM wbs_templates"),
                ("pending_approvals", "SELECT count(*) FROM wbs_approval_requests WHERE status = 'Pending'"),
                ("preview_imports", "SELECT count(*) FROM wbs_import_jobs WHERE status = 'Preview'"),
                ("locked_baselines", "SELECT count(*) FROM wbs_project_baselines WHERE status = 'Locked'"),
            )
        }
        status_rows = await connection.fetch(
            "SELECT status, count(*) AS count FROM wbs_projects GROUP BY status ORDER BY status"
        )
        projects = await connection.fetch(
            """
            SELECT id, name, owner, status, template_key, start_date,
                   openproject_project_id, created_at, updated_at
            FROM wbs_projects
            ORDER BY updated_at DESC
            LIMIT 100
            """
        )
        approvals = await connection.fetch(
            f"""
            SELECT {APPROVAL_SELECT}
            FROM wbs_approval_requests a
            JOIN wbs_projects p ON p.id = a.project_id
            WHERE a.created_at >= $1 OR a.updated_at >= $1
            ORDER BY a.updated_at DESC
            LIMIT 100
            """,
            period_start,
        )
        imports = await connection.fetch(
            """
            SELECT source_file, template_key, template_name, status, total_rows,
                   accepted_rows, rejected_rows, applied_at, created_at
            FROM wbs_import_jobs
            WHERE created_at >= $1 OR applied_at >= $1
            ORDER BY created_at DESC
            LIMIT 100
            """,
            period_start,
        )
        sync_runs = await connection.fetch(
            f"""
            SELECT {SYNC_RUN_SELECT}
            FROM wbs_sync_runs s
            JOIN wbs_projects p ON p.id = s.project_id
            WHERE s.started_at >= $1
            ORDER BY s.started_at DESC
            LIMIT 100
            """,
            period_start,
        )
        rollups = await fetch_rollups_for_projects(
            connection,
            [row["id"] for row in projects],
            period_end.date() if isinstance(period_end, datetime) else None,
        )
    return {
        **counts,
        "status_rows": status_rows,
        "projects": projects,
        "approvals": approvals,
        "imports": imports,
        "sync_runs": sync_runs,
        "rollups": rollups,
    }


def _progress_row(project: Any, rollup: dict[str, Any]) -> list[Any]:
    """WBS 진행 현황 행 — 작업/완료/지연 건수도 공통 산식의 최하위 작업(leaf) 기준."""
    sm = rollup.get("summary", {})
    return [
        project["name"],
        project["owner"],
        project["status"],
        sm.get("leaf_count", 0),
        sm.get("done_leaf_count", 0),
        sm.get("delayed_leaf_count", 0),
        f"{sm.get('progress_pct', 0)}%",
        f"{sm.get('planned_pct', 0)}%",
        sm.get("spi") or "—",
    ]


def _weekly_report_sheets(
    data: dict[str, Any], period_start: datetime, period_end: datetime
) -> list[tuple[str, list[tuple[str, list[str], list[list[Any]]]]]]:
    """[(시트명, [(표 제목, 헤더, 행들), ...]), ...]"""
    return [
        ("Summary", [
            ("WBS 주간 프로젝트 보고서", ["항목", "값"], [
                ["보고 기간", f"{period_start.date().isoformat()} ~ {period_end.date().isoformat()}"],
                ["생성 시각", period_end],
                ["프로젝트", data["project_count"]],
                ["템플릿", data["template_count"]],
                ["승인 대기", data["pending_approvals"]],
                ["Excel 미반영 Preview", data["preview_imports"]],
                ["Locked Baseline", data["locked_baselines"]],
            ]),
            ("프로젝트 상태 분포", ["상태", "건수"], [[r["status"], r["count"]] for r in data["status_rows"]]),
        ]),
        ("Projects", [
            ("프로젝트 현황", ["프로젝트", "담당", "상태", "템플릿", "시작일", "외부 프로젝트 ID", "생성", "수정"], [
                [r["name"], r["owner"], r["status"], r["template_key"], r["start_date"],
                 r["openproject_project_id"], r["created_at"], r["updated_at"]]
                for r in data["projects"]
            ]),
        ]),
        ("WBS Progress", [
            ("내부 WBS 진행 현황",
             ["프로젝트", "담당", "상태", "작업 항목", "완료", "지연", "진척률(가중)", "계획 진척률", "SPI"],
             [_progress_row(r, data["rollups"].get(str(r["id"]), {})) for r in data["projects"]]),
        ]),
        ("Approvals", [
            ("최근 승인 이력", ["프로젝트", "제목", "상태", "요청자", "검토자", "생성", "결정"], [
                [r["project_name"], r["title"], r["status"], r["requester"], r["reviewer"], r["created_at"], r["decided_at"]]
                for r in data["approvals"]
            ]),
        ]),
        ("Excel Imports", [
            ("최근 Excel 반영", ["파일", "템플릿", "상태", "전체", "정상", "오류", "생성", "반영"], [
                [r["source_file"], r["template_name"] or r["template_key"], r["status"], r["total_rows"],
                 r["accepted_rows"], r["rejected_rows"], r["created_at"], r["applied_at"]]
                for r in data["imports"]
            ]),
        ]),
        ("Baseline Runs", [
            ("최근 기준선 반영 실행", ["프로젝트", "모드", "상태", "전체", "생성", "외부 프로젝트 ID", "시작", "완료"], [
                [r["project_name"], r["mode"], r["status"], r["total_rows"], r["created_work_packages"],
                 r["openproject_project_id"], r["started_at"], r["completed_at"]]
                for r in data["sync_runs"]
            ]),
        ]),
    ]


async def build_weekly_report_workbook(
    pool: asyncpg.Pool,
    period_start: datetime,
    period_end: datetime,
) -> tuple[bytes, dict[str, Any]]:
    data = await _fetch_weekly_report_data(pool, period_start, period_end)
    wb = Workbook()
    for index, (sheet_name, tables) in enumerate(_weekly_report_sheets(data, period_start, period_end)):
        worksheet = wb.active if index == 0 else wb.create_sheet(sheet_name)
        worksheet.title = sheet_name
        for title, headers, rows in tables:
            append_report_table(worksheet, title, headers, rows)
        autosize_worksheet(worksheet)

    buffer = BytesIO()
    wb.save(buffer)
    metadata = {
        "project_count": data["project_count"],
        "pending_approvals": data["pending_approvals"],
        "preview_imports": data["preview_imports"],
        "sync_runs": len(data["sync_runs"]),
        "wbs_progress_projects": len(data["projects"]),
        "approvals": len(data["approvals"]),
        "imports": len(data["imports"]),
        "data_source": "internal_wbs",
    }
    return buffer.getvalue(), metadata
