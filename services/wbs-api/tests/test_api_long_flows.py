"""4단계(긴 함수 분해) 대상 기능의 동작 고정 테스트 — 분해 전후 결과가 같아야 한다.

대상: PM 엔진 동기화(sync_project_to_engine), 주간 Excel 보고서, 운영 점검(operations/health)
"""
import io
import unittest

from tests.harness import SAMPLE_WBS, env

EXPECTED_CHECK_IDS = {
    "postgresql", "schema", "template_baseline", "portfolio_seed", "approval_queue", "excel_preview_queue",
    "openproject_preflight", "sync_audit", "baseline_lock", "access_control", "settings_registry", "audit_log",
    "account_lockout", "template_versions", "project_wbs_storage", "weekly_report_scheduler", "cors_policy",
    "backup_rehearsal", "metrics",
}


def setUpModule():
    env()


class OperationsHealthTests(unittest.TestCase):
    def test_health_checks_shape(self):  # noqa: D401
        r = env().get("/api/operations/health")
        self.assertEqual(r.status, 200, r)
        ids = {c["key"] for c in r.body["checks"]}
        self.assertEqual(ids, EXPECTED_CHECK_IDS)
        schema = next(c for c in r.body["checks"] if c["key"] == "schema")
        self.assertEqual(schema["status"], "pass", schema)  # 신규 설치 DB 에 필수 테이블 누락 없음
        s = r.body["summary"]
        self.assertEqual(s["checks"], len(EXPECTED_CHECK_IDS))
        self.assertEqual(s["checks"], s["failures"] + s["warnings"] + s["passes"])
        self.assertIn(r.body["status"], {"stable", "watch", "critical"})

    def test_viewer_forbidden(self):
        env().ensure_user("viewer_ops", "Viewer Ops", "viewer")
        self.assertEqual(env().get("/api/operations/health", as_="viewer_ops").status, 403)


class WeeklyReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        e = env()
        cls.name = "[TEST] 주간보고 진척"
        cls.pid = e.create_project(cls.name)
        assert e.save_wbs(cls.pid, SAMPLE_WBS).status == 200
        e.patch(f"/api/projects/{cls.pid}/work-items/P.1.1", {"progress": 100})
        e.patch(f"/api/projects/{cls.pid}/work-items/P.1.2", {"progress": 50})

    def test_workbook_sheets_and_progress_row(self):
        """정기 보고 스케줄(weekly-pmo) 실행 → 산출물 Excel 내용 검증 (build_weekly_report_workbook)."""
        from openpyxl import load_workbook

        e = env()
        run = e.post("/api/report-schedules/weekly-pmo/run", {"send_email": False, "triggered_by": "test"})
        self.assertEqual(run.status, 201, run)
        self.assertEqual(run.body["run"]["status"], "Generated", run)
        art = e.get(f"/api/report-runs/{run.body['run']['id']}/artifact")
        self.assertEqual(art.status, 200, art)
        wb = load_workbook(io.BytesIO(art.body))
        self.assertEqual(wb.sheetnames, ["Summary", "Projects", "WBS Progress", "Approvals", "Excel Imports", "Baseline Runs"])
        rows = [[c.value for c in row] for row in wb["WBS Progress"].iter_rows()]
        header = next(row for row in rows if row and row[0] == "프로젝트")
        mine = next(row for row in rows if row and row[0] == self.name)
        rec = dict(zip(header, mine))
        self.assertEqual(rec["진척률(가중)"], "30.0%")
        self.assertEqual(rec["계획 진척률"], "40.0%")
        self.assertEqual(str(rec["SPI"]), "0.75")  # 보고서 셀은 문자열로 기록됨
        # 작업/완료/지연 건수도 공통 산식의 최하위 작업(leaf) 기준
        self.assertEqual((str(rec["작업 항목"]), str(rec["완료"]), str(rec["지연"])), ("5", "1", "1"))


class PmEngineSyncTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        e = env()
        cls.pid = e.create_project("[TEST] 동기화")
        assert e.save_wbs(cls.pid, SAMPLE_WBS).status == 200

    def test_1_dry_run(self):
        r = env().post(f"/api/projects/{self.pid}/sync", {"dry_run": True})
        self.assertEqual(r.status, 200, r)
        self.assertEqual(r.body["status"], "DryRun")
        self.assertIn("rows", r.body)
        self.assertEqual(r.body["audit"]["status"], "DryRun")

    def test_2_actual_sync_requires_approved_project(self):
        r = env().post(f"/api/projects/{self.pid}/sync", {"dry_run": False})
        self.assertEqual(r.status, 409, r)

    def test_3_actual_sync_with_mock_adapter(self):
        e = env()
        ap = e.post("/api/approvals", {"project_id": self.pid, "title": "동기화 승인"})
        self.assertEqual(ap.status, 201, ap)
        if ap.body.get("status") != "Approved":
            e.post(f"/api/approvals/{ap.body['id']}/approve", {"reviewer": "PMO Lead"})
        r = e.post(f"/api/projects/{self.pid}/sync", {"dry_run": False})
        self.assertEqual(r.status, 200, r)
        self.assertEqual(r.body["status"], "Synced")
        self.assertEqual(r.body["project"]["status"], "Synced")
        self.assertEqual(r.body["summary"]["created_work_packages"], len(r.body["created_work_packages"]))
        self.assertGreater(r.body["summary"]["created_work_packages"], 0)
        runs = e.get(f"/api/projects/{self.pid}/sync-runs").body
        runs = runs.get("runs", runs) if isinstance(runs, dict) else runs
        self.assertIn("Synced", [x["status"] for x in runs])

    def test_4_resync_after_synced_is_blocked(self):
        r = env().post(f"/api/projects/{self.pid}/sync", {"dry_run": False})
        self.assertEqual(r.status, 409, r)


if __name__ == "__main__":
    unittest.main()
